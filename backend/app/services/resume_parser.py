"""
Resume text extraction and parsing service.
Supports PDF (via PyMuPDF) and DOCX (via python-docx).
"""

import re
from pathlib import Path
from typing import Optional
import pymupdf as fitz  # PyMuPDF (renamed from 'fitz'; kept alias for API compatibility)
from docx import Document


def extract_text_from_pdf(file_path: str) -> tuple[str, dict]:
    """Extract text from a PDF file."""
    doc = fitz.open(file_path)
    pages_text = []
    images_found = 0

    for page_num in range(len(doc)):
        page = doc[page_num]
        text = page.get_text("text")
        pages_text.append(text)
        images_found += len(page.get_images(full=True))

    full_text = "\n".join(pages_text)
    metadata = {
        "page_count": len(doc),
        "images_found": images_found,
        "total_chars": len(full_text),
        "extraction_method": "pymupdf",
    }
    doc.close()
    return full_text, metadata


def extract_text_from_docx(file_path: str) -> tuple[str, dict]:
    """Extract text from a DOCX file."""
    doc = Document(file_path)
    paragraphs = []
    tables_found = 0

    for para in doc.paragraphs:
        if para.text.strip():
            paragraphs.append(para.text)

    for table in doc.tables:
        tables_found += 1
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                paragraphs.append(row_text)

    full_text = "\n".join(paragraphs)
    metadata = {
        "paragraph_count": len(doc.paragraphs),
        "tables_found": tables_found,
        "total_chars": len(full_text),
        "extraction_method": "python-docx",
    }
    return full_text, metadata


def extract_resume_text(file_path: str) -> tuple[str, dict]:
    """Extract text from a resume file (PDF or DOCX)."""
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"File not found: {file_path}")

    ext = path.suffix.lower()
    if ext == ".pdf":
        text, meta = extract_text_from_pdf(file_path)
    elif ext == ".docx":
        text, meta = extract_text_from_docx(file_path)
    else:
        raise ValueError(f"Unsupported file type: {ext}")

    if not text.strip():
        raise ValueError("Could not extract readable text from the file.")

    return text, meta


def extract_sections(text: str) -> dict:
    """Identify and extract common resume sections."""
    section_patterns = {
        "summary": r"(?i)^(summary|profile|objective|about\s*me|professional\s*summary)\s*$",
        "education": r"(?i)^(education|academic|qualification|degree)\s*$",
        "experience": r"(?i)^(experience|work\s*experience|employment|work\s*history)\s*$",
        "skills": r"(?i)^(skills|technical\s*skills|competencies|technologies|tech\s*stack)\s*$",
        "projects": r"(?i)^(projects|personal\s*projects|key\s*projects)\s*$",
        "certifications": r"(?i)^(certifications?|certificates?|licenses?|credentials)\s*$",
        "achievements": r"(?i)^(achievements?|awards?|honors?|accomplishments?)\s*$",
        "languages": r"(?i)^(languages?|language\s*proficiency)\s*$",
    }

    lines = text.split("\n")
    sections = {}
    current_section = None
    current_content = []

    for line in lines:
        stripped = line.strip()
        if not stripped:
            current_content.append("")
            continue

        matched = False
        for section_name, pattern in section_patterns.items():
            if re.match(pattern, stripped):
                if current_section:
                    sections[current_section] = "\n".join(current_content).strip()
                current_section = section_name
                current_content = []
                matched = True
                break
        if not matched:
            current_content.append(stripped)

    if current_section:
        sections[current_section] = "\n".join(current_content).strip()

    return sections


def extract_contact_info(text: str) -> dict:
    """Extract contact information from resume text."""
    contact = {"email": None, "phone": None, "linkedin": None, "github": None, "name": None}

    email_match = re.search(r'[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}', text)
    if email_match:
        contact["email"] = email_match.group(0)

    phone_match = re.search(r'(?:\+?1[-.\s]?)?\(?\d{3}\)?[-.\s]?\d{3}[-.\s]?\d{4}', text)
    if phone_match:
        contact["phone"] = phone_match.group(0).strip()

    linkedin_match = re.search(r'linkedin\.com/in/[a-zA-Z0-9_-]+', text, re.IGNORECASE)
    if linkedin_match:
        contact["linkedin"] = "https://" + linkedin_match.group(0)

    github_match = re.search(r'github\.com/[a-zA-Z0-9_-]+', text, re.IGNORECASE)
    if github_match:
        contact["github"] = "https://" + github_match.group(0)

    lines = [l.strip() for l in text.split("\n") if l.strip()]
    for line in lines[:5]:
        words = line.split()
        if 1 < len(words) <= 4 and all(w.replace(".", "").replace("'", "").isalpha() for w in words):
            skip_words = ["email", "phone", "linkedin", "github", "@", "summary", "experience", "education", "skills"]
            if not any(kw in line.lower() for kw in skip_words):
                contact["name"] = line
                break

    return contact


def extract_skills_from_text(text: str) -> list[str]:
    """Extract potential skill keywords from text using pattern matching."""
    skills = set()
    tech_patterns = [
        r'\b(Python|Java|JavaScript|TypeScript|C\+\+|C#|Ruby|Go|Rust|PHP|Swift|Kotlin|Scala|R|MATLAB|SQL|NoSQL)\b',
        r'\b(React|Angular|Vue\.?js|Next\.?js|Node\.?js|Django|Flask|FastAPI|Spring|Express|Rails|Laravel|ASP\.NET)\b',
        r'\b(AWS|Azure|GCP|Docker|Kubernetes|Jenkins|Git|GitHub|GitLab|CI.CD|Terraform|Ansible)\b',
        r'\b(MySQL|PostgreSQL|MongoDB|Redis|Elasticsearch|DynamoDB|SQLite|Oracle)\b',
        r'\b(REST|RESTful|GraphQL|gRPC|WebSocket|Microservices)\b',
        r'\b(TensorFlow|PyTorch|Scikit-learn|Pandas|NumPy|Machine Learning|Deep Learning|NLP)\b',
        r'\b(Agile|Scrum|Kanban|JIRA|Confluence)\b',
        r'\b(HTML|CSS|SASS|Tailwind|Bootstrap)\b',
        r'\b(Linux|Unix|Bash|Shell|PowerShell)\b',
    ]
    for pattern in tech_patterns:
        matches = re.findall(pattern, text, re.IGNORECASE)
        for match in matches:
            skills.add(match.strip())
    return sorted(skills)

